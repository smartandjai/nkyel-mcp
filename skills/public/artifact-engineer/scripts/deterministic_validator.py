"""Deterministic Validator for Ñkyel Studio Deliverables: fast, non-LLM checks before Vision QA."""

from __future__ import annotations

import os
import zipfile
from pathlib import Path
from typing import Any
import pypdf
import openpyxl

from document_spec import DocumentSpec


class ValidationResult:
    def __init__(self):
        self.valid: bool = True
        self.errors: list[str] = []
        self.warnings: list[str] = []
        self.stats: dict[str, Any] = {}

    def add_error(self, msg: str):
        self.valid = False
        self.errors.append(msg)

    def add_warning(self, msg: str):
        self.warnings.append(msg)

    def to_dict(self) -> dict[str, Any]:
        return {
            "valid": self.valid,
            "errors": self.errors,
            "warnings": self.warnings,
            "stats": self.stats,
        }


def validate_pdf(pdf_path: Path, doc: DocumentSpec) -> ValidationResult:
    res = ValidationResult()
    if not pdf_path.exists():
        res.add_error(f"PDF file does not exist: {pdf_path}")
        return res
    
    size = os.path.getsize(pdf_path)
    res.stats["file_size"] = size
    if size < 1000:
        res.add_error(f"PDF file too small ({size} bytes), possible corruption.")
        return res

    with open(pdf_path, "rb") as f:
        header = f.read(5)
        if not header.startswith(b"%PDF-"):
            res.add_error("Invalid PDF header magic bytes.")
            return res

    try:
        reader = pypdf.PdfReader(str(pdf_path))
        page_count = len(reader.pages)
        res.stats["page_count"] = page_count
        if page_count == 0:
            res.add_error("PDF contains 0 pages.")
        
        # Verify text extractability
        has_text = False
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if len(text.strip()) > 10:
                has_text = True
            else:
                res.add_warning(f"Page {i+1} has very little or no extractable text.")
        
        if not has_text:
            res.add_error("PDF contains no readable text.")

    except Exception as e:
        res.add_error(f"pypdf validation failed: {str(e)}")

    return res


def validate_docx(docx_path: Path, doc: DocumentSpec) -> ValidationResult:
    res = ValidationResult()
    if not docx_path.exists():
        res.add_error(f"DOCX file does not exist: {docx_path}")
        return res
    
    size = os.path.getsize(docx_path)
    res.stats["file_size"] = size
    if size < 2000:
        res.add_error(f"DOCX file too small ({size} bytes).")
        return res

    try:
        with zipfile.ZipFile(docx_path, "r") as z:
            file_list = z.namelist()
            if "word/document.xml" not in file_list:
                res.add_error("DOCX archive is missing word/document.xml.")
                return res
            doc_xml = z.read("word/document.xml").decode("utf-8")
            if doc.title and doc.title not in doc_xml:
                res.add_warning("Document title not found in word/document.xml.")
            res.stats["xml_length"] = len(doc_xml)
    except Exception as e:
        res.add_error(f"DOCX zip validation error: {str(e)}")

    return res


def validate_pptx(pptx_path: Path, doc: DocumentSpec) -> ValidationResult:
    res = ValidationResult()
    if not pptx_path.exists():
        res.add_error(f"PPTX file does not exist: {pptx_path}")
        return res
    
    size = os.path.getsize(pptx_path)
    res.stats["file_size"] = size
    if size < 2000:
        res.add_error(f"PPTX file too small ({size} bytes).")
        return res

    try:
        with zipfile.ZipFile(pptx_path, "r") as z:
            file_list = z.namelist()
            if "ppt/presentation.xml" not in file_list:
                res.add_error("PPTX archive is missing ppt/presentation.xml.")
                return res
            
            slide_entries = [f for f in file_list if f.startswith("ppt/slides/slide") and f.endswith(".xml")]
            res.stats["slide_count"] = len(slide_entries)
            if len(doc.slides) > 0 and len(slide_entries) < len(doc.slides):
                res.add_error(f"Expected {len(doc.slides)} slides, found {len(slide_entries)}.")
    except Exception as e:
        res.add_error(f"PPTX zip validation error: {str(e)}")

    return res


def validate_xlsx(xlsx_path: Path, doc: DocumentSpec) -> ValidationResult:
    res = ValidationResult()
    if not xlsx_path.exists():
        res.add_error(f"XLSX file does not exist: {xlsx_path}")
        return res
    
    size = os.path.getsize(xlsx_path)
    res.stats["file_size"] = size
    if size < 1000:
        res.add_error(f"XLSX file too small ({size} bytes).")
        return res

    try:
        wb = openpyxl.load_workbook(str(xlsx_path), data_only=False)
        res.stats["sheet_count"] = len(wb.worksheets)
        if len(doc.sheets) > 0 and len(wb.worksheets) < len(doc.sheets):
            res.add_error(f"Expected {len(doc.sheets)} sheets, found {len(wb.worksheets)}.")

        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for cell in row:
                    if isinstance(cell, str):
                        for bad_formula in ["#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A"]:
                            if bad_formula in cell:
                                res.add_error(f"Spreadsheet formula error found in sheet '{ws.title}': {bad_formula}")
    except Exception as e:
        res.add_error(f"XLSX validation error: {str(e)}")

    return res

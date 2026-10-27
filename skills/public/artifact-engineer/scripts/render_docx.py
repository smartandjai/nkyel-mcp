"""DOCX Renderer using python-docx: produces true editable Microsoft Word documents."""

from __future__ import annotations

from pathlib import Path
from typing import Any
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from document_spec import DocumentSpec, TableData


def _hex_to_rgb(hex_str: str) -> RGBColor:
    hex_str = hex_str.lstrip("#")
    r, g, b = int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16)
    return RGBColor(r, g, b)


def _set_cell_background(cell, hex_color: str):
    """Set background shading for a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color.lstrip("#")}"/>')
    tc_pr.append(shd)


def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner cell padding."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def render_docx(doc_spec: DocumentSpec, output_dir: Path) -> dict[str, Any]:
    """Generate professional editable DOCX deliverable from DocumentSpec."""
    output_dir.mkdir(parents=True, exist_ok=True)
    docx_path = output_dir / f"{doc_spec.document_id}.docx"
    
    doc = Document()
    theme = doc_spec.theme
    
    # 1. Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header & Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"{doc_spec.title} · {doc_spec.organization}")
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = _hex_to_rgb(theme.muted_color)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run(f"Confidentiel · {doc_spec.date} · Généré par Ñkyel Studio")
        frun.font.size = Pt(8)
        frun.font.color.rgb = _hex_to_rgb(theme.muted_color)

    # 2. Cover / Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(36)
    title_p.paragraph_format.space_after = Pt(8)
    title_run = title_p.add_run(doc_spec.title)
    title_run.font.name = theme.font_family_headline
    title_run.font.size = Pt(26)
    title_run.font.bold = True
    title_run.font.color.rgb = _hex_to_rgb(theme.primary_color)

    if doc_spec.subtitle:
        sub_p = doc.add_paragraph()
        sub_p.paragraph_format.space_after = Pt(18)
        sub_run = sub_p.add_run(doc_spec.subtitle)
        sub_run.font.size = Pt(13)
        sub_run.font.color.rgb = _hex_to_rgb(theme.secondary_color)

    # Metadata banner
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(24)
    meta_run = meta_p.add_run(f"Auteur : {doc_spec.author}  |  Organisation : {doc_spec.organization}  |  Date : {doc_spec.date}")
    meta_run.font.size = Pt(9.5)
    meta_run.font.color.rgb = _hex_to_rgb(theme.muted_color)

    # 3. Executive Summary
    if doc_spec.executive_summary:
        exec_table = doc.add_table(rows=1, cols=1)
        exec_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = exec_table.cell(0, 0)
        _set_cell_background(cell, "#f8fafc")
        _set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        
        cp = cell.paragraphs[0]
        crun_title = cp.add_run("RÉSUMÉ EXÉCUTIF\n")
        crun_title.font.bold = True
        crun_title.font.size = Pt(11)
        crun_title.font.color.rgb = _hex_to_rgb(theme.primary_color)
        
        crun_text = cp.add_run(doc_spec.executive_summary)
        crun_text.font.size = Pt(10)
        crun_text.font.color.rgb = _hex_to_rgb(theme.text_color)
        
        doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 4. Sections
    for sec in doc_spec.sections:
        if sec.page_break_before:
            doc.add_page_break()

        h_p = doc.add_paragraph()
        h_p.paragraph_format.space_before = Pt(18)
        h_p.paragraph_format.space_after = Pt(6)
        h_p.paragraph_format.keep_with_next = True
        h_run = h_p.add_run(sec.title)
        h_run.font.bold = True
        h_run.font.size = Pt(16 if sec.level == 1 else 13)
        h_run.font.color.rgb = _hex_to_rgb(theme.primary_color)

        # Paragraphs
        for p_block in sec.paragraphs:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            
            if p_block.style == "callout":
                p.paragraph_format.left_indent = Inches(0.3)
                run = p.add_run(f"« {p_block.text} »")
                run.font.italic = True
                run.font.color.rgb = _hex_to_rgb(theme.secondary_color)
            elif p_block.style == "bullet":
                p.style = "List Bullet"
                run = p.add_run(p_block.text)
            else:
                run = p.add_run(p_block.text)
                run.font.color.rgb = _hex_to_rgb(theme.text_color)
                run.font.size = Pt(10.5)

        # Tables
        for table_data in sec.tables:
            if table_data.title:
                tp = doc.add_paragraph()
                tp.paragraph_format.space_before = Pt(8)
                tp.paragraph_format.space_after = Pt(4)
                trun = tp.add_run(table_data.title)
                trun.font.bold = True
                trun.font.size = Pt(10.5)
                trun.font.color.rgb = _hex_to_rgb(theme.primary_color)

            cols = table_data.columns
            rows = table_data.rows
            if not cols:
                continue

            t = doc.add_table(rows=len(rows) + 1, cols=len(cols))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            # Header Row
            hdr_cells = t.rows[0].cells
            for idx, col in enumerate(cols):
                hdr_cells[idx].text = col.header
                _set_cell_background(hdr_cells[idx], theme.primary_color)
                _set_cell_margins(hdr_cells[idx], top=100, bottom=100, left=120, right=120)
                for run in hdr_cells[idx].paragraphs[0].runs:
                    run.font.bold = True
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(255, 255, 255)

            # Data Rows
            for r_idx, row in enumerate(rows):
                row_cells = t.rows[r_idx + 1].cells
                bg_color = "#f8fafc" if r_idx % 2 == 1 else "#ffffff"
                for c_idx, col in enumerate(cols):
                    val = str(row.get(col.key, ""))
                    row_cells[c_idx].text = val
                    _set_cell_background(row_cells[c_idx], bg_color)
                    _set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=120, right=120)
                    for run in row_cells[c_idx].paragraphs[0].runs:
                        run.font.size = Pt(9)
                        run.font.color.rgb = _hex_to_rgb(theme.text_color)
            
            doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 5. Citations
    if doc_spec.citations:
        doc.add_paragraph().paragraph_format.space_before = Pt(16)
        ref_h = doc.add_paragraph()
        ref_h.paragraph_format.keep_with_next = True
        rrun = ref_h.add_run("Références & Sources")
        rrun.font.bold = True
        rrun.font.size = Pt(14)
        rrun.font.color.rgb = _hex_to_rgb(theme.primary_color)
        
        for idx, cit in enumerate(doc_spec.citations, 1):
            cp = doc.add_paragraph()
            cp.paragraph_format.space_after = Pt(3)
            crun1 = cp.add_run(f"[{idx}] {cit.title}")
            crun1.font.bold = True
            crun1.font.size = Pt(9)
            if cit.notes:
                crun2 = cp.add_run(f" — {cit.notes}")
                crun2.font.size = Pt(9)
            if cit.url:
                crun3 = cp.add_run(f" ({cit.url})")
                crun3.font.size = Pt(8.5)
                crun3.font.color.rgb = _hex_to_rgb(theme.secondary_color)

    doc.save(str(docx_path))
    return {
        "docx_path": str(docx_path),
        "file_size": docx_path.stat().st_size,
    }

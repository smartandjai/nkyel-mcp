"""Editable PPTX Renderer using python-pptx: produces true native 16:9 PowerPoint slides."""

from __future__ import annotations

from pathlib import Path
from typing import Any
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from document_spec import DocumentSpec, SlideContent, TableData


def _hex_to_rgb(hex_str: str) -> RGBColor:
    hex_str = hex_str.lstrip("#")
    r, g, b = int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16)
    return RGBColor(r, g, b)


def render_pptx(doc_spec: DocumentSpec, output_dir: Path) -> dict[str, Any]:
    """Generate professional native editable 16:9 PPTX deliverable."""
    output_dir.mkdir(parents=True, exist_ok=True)
    pptx_path = output_dir / f"{doc_spec.document_id}.pptx"

    prs = Presentation()
    # Standard 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    theme = doc_spec.theme
    primary_color = _hex_to_rgb(theme.primary_color)
    secondary_color = _hex_to_rgb(theme.secondary_color)
    text_color = _hex_to_rgb(theme.text_color)
    muted_color = _hex_to_rgb(theme.muted_color)
    bg_color = _hex_to_rgb(theme.bg_color)

    blank_layout = prs.slide_layouts[6]  # Blank slide

    for slide_data in doc_spec.slides:
        slide = prs.slides.add_slide(blank_layout)

        # Background Fill
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = bg_color

        # Speaker notes
        if slide_data.speaker_notes:
            notes_slide = slide.notes_slide
            text_frame = notes_slide.notes_text_frame
            text_frame.text = slide_data.speaker_notes

        # Layout Dispatch
        if slide_data.layout == "title":
            _render_title_slide(slide, slide_data, doc_spec, primary_color, secondary_color, muted_color)
        elif slide_data.layout == "metric_highlight":
            _render_metrics_slide(slide, slide_data, primary_color, secondary_color, text_color, muted_color)
        elif slide_data.layout == "table" and slide_data.table:
            _render_table_slide(slide, slide_data, primary_color, secondary_color, text_color, muted_color)
        else:
            _render_content_slide(slide, slide_data, primary_color, secondary_color, text_color, muted_color)

    prs.save(str(pptx_path))
    return {
        "pptx_path": str(pptx_path),
        "slide_count": len(doc_spec.slides),
        "file_size": pptx_path.stat().st_size,
    }


def _render_title_slide(slide, slide_data: SlideContent, doc_spec: DocumentSpec, p_color, s_color, m_color):
    """Render a premium cover/title slide."""
    # Top decorative colored bar
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = s_color
    top_bar.line.fill.background()

    # Title & Subtitle text frame
    tx_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.333), Inches(3.0))
    tf = tx_box.text_frame
    tf.word_wrap = True

    p1 = tf.paragraphs[0]
    p1.text = slide_data.title
    p1.font.bold = True
    p1.font.size = Pt(40)
    p1.font.color.rgb = p_color
    p1.space_after = Pt(14)

    if slide_data.subtitle:
        p2 = tf.add_paragraph()
        p2.text = slide_data.subtitle
        p2.font.size = Pt(20)
        p2.font.color.rgb = s_color
        p2.space_after = Pt(24)

    p3 = tf.add_paragraph()
    p3.text = f"{doc_spec.organization} · {doc_spec.author} · {doc_spec.date}"
    p3.font.size = Pt(12)
    p3.font.color.rgb = m_color


def _render_content_slide(slide, slide_data: SlideContent, p_color, s_color, t_color, m_color):
    """Render a standard content slide with bullets."""
    # Header area
    _render_slide_header(slide, slide_data.title, p_color, s_color, m_color)

    # Bullets container
    tx_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.8), Inches(4.5))
    tf = tx_box.text_frame
    tf.word_wrap = True

    for idx, bullet in enumerate(slide_data.bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"•  {bullet}"
        p.font.size = Pt(18)
        p.font.color.rgb = t_color
        p.space_after = Pt(16)
        p.level = 0


def _render_metrics_slide(slide, slide_data: SlideContent, p_color, s_color, t_color, m_color):
    """Render KPI metrics in balanced cards."""
    _render_slide_header(slide, slide_data.title, p_color, s_color, m_color)

    metrics = slide_data.metrics or []
    count = len(metrics)
    if not count:
        return

    card_width = min(2.8, (10.8 - (count - 1) * 0.4) / count)
    left_start = Inches(1.2)
    gap = Inches(0.4)

    for idx, m in enumerate(metrics):
        card_left = left_start + idx * (Inches(card_width) + gap)
        # Card shape
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_left, Inches(2.4), Inches(card_width), Inches(3.2))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(248, 250, 252)
        card.line.color.rgb = RGBColor(226, 232, 240)

        tf = card.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        p1.text = "\n" + m.get("value", "")
        p1.font.bold = True
        p1.font.size = Pt(32)
        p1.font.color.rgb = s_color
        p1.space_after = Pt(12)

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.text = m.get("label", "")
        p2.font.bold = True
        p2.font.size = Pt(13)
        p2.font.color.rgb = p_color


def _render_table_slide(slide, slide_data: SlideContent, p_color, s_color, t_color, m_color):
    """Render a native PowerPoint table."""
    _render_slide_header(slide, slide_data.title, p_color, s_color, m_color)
    table_data = slide_data.table
    if not table_data or not table_data.columns:
        return

    rows = table_data.rows
    cols = table_data.columns

    table_shape = slide.shapes.add_table(len(rows) + 1, len(cols), Inches(1.2), Inches(2.0), Inches(10.8), Inches(4.5))
    table = table_shape.table

    # Headers
    for c_idx, col in enumerate(cols):
        cell = table.cell(0, c_idx)
        cell.text = col.header
        cell.fill.solid()
        cell.fill.fore_color.rgb = p_color
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(12)
            p.font.color.rgb = RGBColor(255, 255, 255)

    # Data
    for r_idx, row in enumerate(rows):
        bg = RGBColor(248, 250, 252) if r_idx % 2 == 1 else RGBColor(255, 255, 255)
        for c_idx, col in enumerate(cols):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = str(row.get(col.key, ""))
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(11)
                p.font.color.rgb = t_color


def _render_slide_header(slide, title: str, p_color, s_color, m_color):
    """Standard header line for presentation slides."""
    tx_box = slide.shapes.add_textbox(Inches(1.2), Inches(0.7), Inches(10.8), Inches(1.0))
    tf = tx_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.bold = True
    p.font.size = Pt(26)
    p.font.color.rgb = p_color

    # Thin accent divider
    divider = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.65), Inches(10.8), Inches(0.04))
    divider.fill.solid()
    divider.fill.fore_color.rgb = s_color
    divider.line.fill.background()

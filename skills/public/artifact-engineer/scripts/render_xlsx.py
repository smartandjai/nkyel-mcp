"""XLSX Renderer using openpyxl: creates professional workbooks with formulas, styling, freeze panes, and auto-fit."""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from document_spec import DocumentSpec, SheetData, SheetColumn


def _hex_to_argb(hex_str: str) -> str:
    """Convert hex color #RRGGBB to openpyxl ARGB format (AARRGGBB)."""
    hex_clean = hex_str.lstrip("#")
    if len(hex_clean) == 6:
        return f"FF{hex_clean.upper()}"
    return "FF1E293B"


def render_xlsx(doc: DocumentSpec, output_dir: Path) -> dict[str, Any]:
    """Generate a clean, professional XLSX workbook from DocumentSpec."""
    output_dir.mkdir(parents=True, exist_ok=True)
    xlsx_path = output_dir / f"{doc.document_id}.xlsx"
    
    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active
    
    theme = doc.theme
    primary_argb = _hex_to_argb(theme.primary_color)
    accent_argb = _hex_to_argb(theme.accent_color)
    
    # Styles
    header_fill = PatternFill(start_color=primary_argb, end_color=primary_argb, fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    total_fill = PatternFill(start_color="FFF1F5F9", end_color="FFF1F5F9", fill_type="solid")
    total_font = Font(name="Calibri", size=10, bold=True, color="FF0F172A")
    
    data_font = Font(name="Calibri", size=10, color="FF0F172A")
    thin_border_side = Side(border_style="thin", color="FFE2E8F0")
    data_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    double_bottom_side = Side(border_style="double", color="FF1E293B")
    top_thin_side = Side(border_style="thin", color="FF1E293B")
    total_border = Border(left=thin_border_side, right=thin_border_side, top=top_thin_side, bottom=double_bottom_side)

    sheets_to_process = doc.sheets
    if not sheets_to_process:
        # Default placeholder sheet if none provided
        ws = wb.create_sheet(title="Résumé")
        ws.append(["Titre", doc.title])
        ws.append(["Auteur", doc.author])
        ws.append(["Organisation", doc.organization])
        ws.append(["Date", doc.date])
    else:
        for sheet_idx, sheet_data in enumerate(sheets_to_process):
            title = sheet_data.name[:31]  # Excel 31-char limit
            ws = wb.create_sheet(title=title)
            ws.views.sheetView[0].showGridLines = True
            
            # Header Row
            headers = [col.header for col in sheet_data.columns]
            ws.append(headers)
            ws.row_dimensions[1].height = 28
            
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_align
                cell.border = data_border
            
            # Data Rows
            start_data_row = 2
            end_data_row = start_data_row + len(sheet_data.rows) - 1
            
            for row_idx, row_dict in enumerate(sheet_data.rows, start=start_data_row):
                ws.row_dimensions[row_idx].height = 20
                for col_idx, col in enumerate(sheet_data.columns, start=1):
                    raw_val = row_dict.get(col.key, "")
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.font = data_font
                    cell.border = data_border
                    
                    # Number / Type conversion
                    num_fmt = col.number_format.lower()
                    if isinstance(raw_val, (int, float)):
                        cell.value = raw_val
                        if "currency" in num_fmt or "€" in num_fmt or "$" in num_fmt:
                            cell.number_format = '#,##0.00 €'
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                        elif "percent" in num_fmt or "%" in num_fmt:
                            cell.number_format = '0.0%'
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                        else:
                            cell.number_format = '#,##0' if isinstance(raw_val, int) else '#,##0.00'
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.value = raw_val
                        cell.alignment = Alignment(horizontal="left", vertical="center")
            
            # Optional Total Row
            if sheet_data.has_total_row and len(sheet_data.rows) > 0:
                total_row_idx = end_data_row + 1
                ws.row_dimensions[total_row_idx].height = 22
                for col_idx, col in enumerate(sheet_data.columns, start=1):
                    cell = ws.cell(row=total_row_idx, column=col_idx)
                    cell.font = total_font
                    cell.fill = total_fill
                    cell.border = total_border
                    
                    col_letter = get_column_letter(col_idx)
                    if col_idx == 1:
                        cell.value = "Total / Moyenne"
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                    else:
                        # Determine if this column was numeric
                        sample_val = sheet_data.rows[0].get(col.key)
                        if isinstance(sample_val, (int, float)):
                            num_fmt = col.number_format.lower()
                            if "percent" in num_fmt or "%" in num_fmt:
                                cell.value = f"=AVERAGE({col_letter}{start_data_row}:{col_letter}{end_data_row})"
                                cell.number_format = '0.0%'
                            elif "currency" in num_fmt or "€" in num_fmt:
                                cell.value = f"=SUM({col_letter}{start_data_row}:{col_letter}{end_data_row})"
                                cell.number_format = '#,##0.00 €'
                            else:
                                cell.value = f"=SUM({col_letter}{start_data_row}:{col_letter}{end_data_row})"
                                cell.number_format = '#,##0' if isinstance(sample_val, int) else '#,##0.00'
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                        else:
                            cell.value = ""
            
            # Freeze Panes
            if sheet_data.freeze_panes:
                ws.freeze_panes = sheet_data.freeze_panes
            
            # Auto-fit Column Widths
            for col_idx, col in enumerate(sheet_data.columns, start=1):
                col_letter = get_column_letter(col_idx)
                max_len = max(len(str(col.header)), 10)
                for r in sheet_data.rows:
                    v = str(r.get(col.key, ""))
                    if len(v) > max_len:
                        max_len = len(v)
                # Cap max width to 45
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, col.width or 12), 45)
    
    # Remove initial empty sheet
    if default_sheet in wb.worksheets and len(wb.worksheets) > 1:
        wb.remove(default_sheet)
    
    wb.save(str(xlsx_path))
    
    return {
        "xlsx_path": str(xlsx_path),
        "sheet_count": len(wb.worksheets),
        "file_size": os.path.getsize(xlsx_path)
    }
